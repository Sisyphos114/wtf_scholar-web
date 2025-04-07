import markdown
from docx import Document
import os
from bs4 import BeautifulSoup  # 用于解析HTML

def convert_md_to_docx(md_filename, docx_filename):
    # 获取当前脚本所在目录
    current_directory = os.path.dirname(os.path.abspath(__file__))
    md_file_path = os.path.join(current_directory, md_filename)
    docx_file_path = os.path.join(current_directory, docx_filename)

    # 读取Markdown文件
    with open(md_file_path, 'r', encoding='utf-8') as md_file:
        md_text = md_file.read()

    # 将Markdown转换为HTML
    html = markdown.markdown(md_text)

    # 创建Word文档
    doc = Document()
    doc.add_heading('Converted Manuscript', 0)

    # 使用BeautifulSoup解析HTML
    soup = BeautifulSoup(html, 'html.parser')

    # 遍历HTML元素并转换为Word格式
    for element in soup:
        if element.name == 'h1':
            doc.add_heading(element.text, level=1)
        elif element.name == 'h2':
            doc.add_heading(element.text, level=2)
        elif element.name == 'p':  # 普通段落
            paragraph = doc.add_paragraph()
            for content in element.contents:
                if isinstance(content, str):
                    paragraph.add_run(content)
                elif content.name == 'strong':  # 加粗
                    paragraph.add_run(content.text).bold = True
                elif content.name == 'em':  # 斜体
                    paragraph.add_run(content.text).italic = True

    # 保存为.docx文件
    doc.save(docx_file_path)
    print(f"文件已保存为: {docx_file_path}")

# 使用示例
convert_md_to_docx('manuscript.md', 'manuscript.docx')
