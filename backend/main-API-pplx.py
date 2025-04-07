from flask import Flask, jsonify, send_file
import os
from dotenv import load_dotenv
import subprocess
from pathlib import Path
from flask import Flask, request, jsonify
from flask_cors import CORS
from docx import Document
import markdown
from bs4 import BeautifulSoup

app = Flask(__name__)
CORS(app)  # 启用 CORS
# 初始化时加载环境变量
load_dotenv()
API_KEY = os.getenv('PPLX_API_KEY')
# 定义 SCRIPT_DIR 为相对于当前文件所在目录的文件夹
SCRIPT_DIR = Path(__file__).parent # 确保脚本在相同目录

@app.route('/run-scripts', methods=['POST'])
def run_all_scripts():
    """执行所有预设脚本的端点，只返回最后一个脚本的执行结果"""
    if not API_KEY:
        return jsonify({"error": "PPLX_API_KEY未配置"}), 500
    
    scripts = [
    "prompt01.py",
    "prompt02pplx.py",
    "prompt03pplx.py",
    "prompt04.py",
    "prompt05.py",
    "prompt06.py"
    ]
    
    result = None  # 初始化为空，后续赋值最后一个脚本的结果
    for script in scripts:
        try:
            script_path = SCRIPT_DIR / script
            if not script_path.exists():
                raise FileNotFoundError(f"脚本 {script} 不存在")
            
            process_result = subprocess.run(
                ['python', str(script_path)],
                check=True,
                text=True,
                capture_output=True,
                encoding='utf-8'
            )
            
            # 将结果保存在变量中，每次覆盖
            result = {
                "script": script,
                "success": True,
                "output": process_result.stdout,
                "error": None
            }
        except subprocess.CalledProcessError as e:
            # 如果脚本出错，停止执行后续脚本，保留出错信息
            result = {
                "script": script,
                "success": False,
                "output": None,
                "error": e.stderr
            }
            break  # 遇到错误时停止执行后续脚本
    
    # Return the final result as a JSON response
    return jsonify(result) if result else jsonify({"error": "没有脚本执行结果"}), 500

# 设置存储文件的路径为当前工作目录下
OUTLINE_FOLDER = os.path.join(os.getcwd())
OUTLINE_FILE = 'outline.txt'  # 固定文件名
os.makedirs(OUTLINE_FOLDER, exist_ok=True)  # 确保文件夹存在

# POST 请求：创建大纲并保存为 outline.txt 文件
@app.route('/api/v1/outlines', methods=['POST'])
def create_outline():
    try:
        # 从请求体中获取大纲数据
        data = request.get_json()

        # 获取论文大纲内容
        outline_text = data.get('outline', '').strip()

        if not outline_text:
            return jsonify({'message': '论文大纲内容不能为空'}), 400

        # 文件的完整路径
        file_path = os.path.join(OUTLINE_FOLDER, OUTLINE_FILE)

        # 将大纲内容保存到 outline.txt 文件中
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(outline_text)

        # 返回文件路径
        return jsonify({'file_path': file_path}), 201

    except Exception as e:
        return jsonify({'message': '创建大纲失败', 'error': str(e)}), 500
    

# 获取当前脚本所在的目录
current_dir = os.path.dirname(os.path.abspath(__file__))
MANUSCRIPT_PATH = os.path.join(current_dir, 'manuscript.md')

# 模拟一个接口获取 manuscript.md 内容
@app.route('/get-manuscript', methods=['GET'])
def get_manuscript():
    try:
        with open(MANUSCRIPT_PATH, 'r', encoding='utf-8') as file:
            content = file.read()
        return content
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# 模拟一个接口来导出 DOCX 文件
@app.route('/export-doc', methods=['POST'])
def export_doc():
    try:
        # 获取修改后的文本内容
        content = request.json.get('content')
        
        if not content:
            return jsonify({"error": "No content provided"}), 400

        # 创建 Word 文档对象
        doc = Document()

        # 将 Markdown 转换为 DOCX
        md_to_docx(doc, content)

        # 保存为临时文件
        temp_path = os.path.join(os.getcwd(), 'temp.docx')
        doc.save(temp_path)

        # 返回文件
        return send_file(temp_path, as_attachment=True)

    except Exception as e:
        return jsonify({"error": str(e)}), 500


def md_to_docx(doc, md_text):
    """将 Markdown 文本转换为 DOCX 并应用格式"""
    # 使用 markdown 库将 Markdown 转换为 HTML
    html_text = markdown.markdown(md_text)
    
    # 使用 BeautifulSoup 解析 HTML
    soup = BeautifulSoup(html_text, 'html.parser')

    # 逐行处理 HTML 内容
    for element in soup.find_all(['p', 'strong', 'em', 'h1', 'h2', 'ul', 'ol']):
        if element.name == 'p':  # 处理段落
            doc.add_paragraph(element.get_text())
        elif element.name == 'strong':  # 处理加粗文本
            doc.add_paragraph().add_run(element.get_text()).bold = True
        elif element.name == 'em':  # 处理斜体文本
            doc.add_paragraph().add_run(element.get_text()).italic = True
        elif element.name == 'h1':  # 处理一级标题
            doc.add_heading(element.get_text(), level=1)
        elif element.name == 'h2':  # 处理二级标题
            doc.add_heading(element.get_text(), level=2)
        elif element.name in ['ul', 'ol']:  # 处理无序/有序列表
            for li in element.find_all('li'):
                doc.add_paragraph(li.get_text(), style='List Bullet' if element.name == 'ul' else 'List Number')

@app.route('/save-manuscript', methods=['POST'])
def save_manuscript():
    try:
        # 获取前端传过来的 content 数据
        content = request.json.get('content')
        
        if not content:
            return jsonify({"success": False, "message": "No content provided"}), 400
        
        # 将 content 保存到 manuscript.md 文件
        with open(MANUSCRIPT_PATH, 'w', encoding='utf-8') as file:
            file.write(content)
        
        # 返回成功的响应
        return jsonify({"success": True, "message": "文档已保存"})

    except Exception as e:
        print(f"Error saving manuscript: {e}")
        return jsonify({"success": False, "message": "保存失败，请稍后重试"}), 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
